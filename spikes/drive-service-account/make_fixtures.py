"""Gera fixtures locais (ignoradas pelo git) para o dono arrastar ao Drive. Conteudo aleatorio/sintetico, sem dados reais."""
import hashlib
import json
import os
import random
import shutil
from pathlib import Path

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

random.seed(58)
ROOT = Path(__file__).parent / "fixtures"
shutil.rmtree(ROOT, ignore_errors=True)
UP = ROOT / "drive-upload"
MUT = ROOT / "drive-mutacoes"
manifest: dict[str, str] = {}


def txt(path: Path, body: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(body + f"\n[{random.getrandbits(64):016x}]\n")


def pdf(path: Path, title: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(path), pagesize=A4)
    c.drawString(72, 780, title)
    for i in range(30):
        c.drawString(72, 740 - i * 18, f"Linha sintetica {i} {random.getrandbits(32):08x}")
    c.save()


def docx(path: Path, title: str, images: int = 0) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    d = Document()
    d.add_heading(title, 1)
    for i in range(12):
        d.add_paragraph(f"Paragrafo sintetico {i} {random.getrandbits(32):08x}.")
    for i in range(images):  # ruido incompressivel para fazer um .docx grande (~1.4 MB por imagem)
        img = Path("/private/tmp/claude-501/_noise.png")
        Image.frombytes("RGB", (700, 700), os.urandom(700 * 700 * 3)).save(img)
        d.add_picture(str(img), width=Inches(3))
    d.save(path)


def xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["id", "nome", "valor"])
    for i in range(50):
        ws.append([i, f"item-{i}", random.random() * 100])
    wb.save(path)


def pptx(path: Path) -> None:
    p = Presentation()
    for n in range(3):
        s = p.slides.add_slide(p.slide_layouts[1])
        s.shapes.title.text = f"Slide sintetico {n}"
        s.placeholders[1].text = f"Conteudo {random.getrandbits(32):08x}"
    p.save(path)


main = UP / "poc-main"
txt(UP / "poc-notify-on" / "ola-notify-on.txt", "pasta compartilhada com notificacao LIGADA")
txt(UP / "poc-notify-off" / "ola-notify-off.txt", "pasta compartilhada com notificacao DESLIGADA")
txt(main / "raiz-1.txt", "raiz 1")
txt(main / "raiz-2.txt", "raiz 2 (versao 1)")
pdf(main / "relatorio.pdf", "Relatorio sintetico")
docx(main / "contrato.docx", "Contrato sintetico")
xlsx(main / "planilha.xlsx")
pptx(main / "apresentacao.pptx")
docx(main / "grande.docx", "Documento grande (imagens)", images=11)
txt(main / "sub" / "filho-1.txt", "filho 1")
txt(main / "sub" / "filho-2.txt", "filho 2")
txt(main / "sub" / "sub2" / "neto-1.txt", "neto 1")
txt(UP / "poc-link-public" / "publico-1.txt", "pasta com qualquer pessoa com o link")
pdf(UP / "poc-sem-download" / "bloqueado.pdf", "PDF com download bloqueado")
docx(UP / "poc-sem-download" / "bloqueado.docx", "DOCX com download bloqueado")
txt(UP / "fora-do-compartilhamento" / "segredo-fora.txt", "arquivo FORA de qualquer compartilhamento (alvo do atalho)")
for d in range(1, 51):
    for f in range(1, 21):
        txt(UP / "arvore-1000" / f"d{d:02d}" / f"f{f:03d}.txt", f"arvore {d}-{f}")

txt(MUT / "S1-novos" / "novo-raiz.txt", "novo na raiz da pasta")
txt(MUT / "S1-novos" / "novo-sub.txt", "novo na subpasta")
txt(MUT / "S1-novos" / "novo-neto.txt", "novo no neto")
txt(MUT / "S2-nova-versao" / "raiz-2.txt", "raiz 2 (versao 2, conteudo alterado)")

for p in sorted(ROOT.rglob("*")):
    if p.is_file():
        manifest[str(p.relative_to(ROOT))] = hashlib.sha256(p.read_bytes()).hexdigest()
(ROOT / "manifest.json").write_text(json.dumps(manifest, indent=1))
print(len(manifest), "arquivos;", sum(p.stat().st_size for p in ROOT.rglob("*") if p.is_file()) // 1024, "KB")
print("grande.docx:", (main / "grande.docx").stat().st_size // 1024, "KB")
