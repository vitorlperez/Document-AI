"""M55: XLSX/PPTX/Sheets/Slides are on by default and bounded so one file cannot exhaust the Worker."""

from io import BytesIO

import pytest
from openpyxl import Workbook
from pptx import Presentation

from app.core.config import Settings
from app.ingestion.extraction import (
    ELIGIBLE_MIME_TYPES,
    eligible_mime_types,
    extract_blocks,
    limits,
)
from app.ingestion.extraction.errors import ExtractionError
from app.ingestion.extraction.mime import PPTX, SHEETS, SLIDES, XLSX
from app.ingestion.extraction.presentation import pptx_blocks
from app.ingestion.extraction.spreadsheet import xlsx_blocks


def workbook_bytes(rows_per_sheet: int, sheets: int = 1, cell: str = "x" * 1000) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for index in range(sheets):
        sheet = workbook.create_sheet(f"Aba{index}")
        sheet.append(["a", "b"])
        for _ in range(rows_per_sheet):
            sheet.append([cell, cell])
    out = BytesIO()
    workbook.save(out)
    return out.getvalue()


def deck_bytes(slides: int) -> bytes:
    presentation = Presentation()
    for number in range(slides):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = f"Slide {number}"
    out = BytesIO()
    presentation.save(out)
    return out.getvalue()


def test_new_formats_are_supported_without_any_flag() -> None:
    assert {XLSX, SHEETS, PPTX, SLIDES, "text/csv"} <= eligible_mime_types(None)
    assert eligible_mime_types(None) == ELIGIBLE_MIME_TYPES
    assert "new_formats_enabled" not in Settings.model_fields


def test_spreadsheet_stops_reading_once_the_document_text_budget_is_spent(monkeypatch) -> None:
    monkeypatch.setattr(limits, "MAX_DOCUMENT_CHARS", 20_000)
    blocks = xlsx_blocks(workbook_bytes(rows_per_sheet=200, sheets=3))

    assert sum(len(block.text) for block in blocks) <= 20_000 + 3_000  # one row of slack
    assert "Conteúdo truncado" in blocks[-1].text
    assert {block.section_path.split(" ›")[0] for block in blocks} == {"Aba0"}  # later sheets never opened


def test_presentation_is_capped_by_slide_count(monkeypatch) -> None:
    monkeypatch.setattr(limits, "MAX_SLIDES", 3)
    blocks = pptx_blocks(deck_bytes(slides=8))

    assert "Conteúdo truncado" in blocks[-1].text
    assert max(block.page_number or 0 for block in blocks) <= 3


@pytest.mark.parametrize("mime", [XLSX, PPTX])
def test_ooxml_formats_have_a_tighter_size_cap_than_other_files(monkeypatch, mime) -> None:
    content = workbook_bytes(rows_per_sheet=5) if mime == XLSX else deck_bytes(slides=2)
    monkeypatch.setattr(limits, "MAX_OOXML_BYTES", len(content) - 1)

    with pytest.raises(ExtractionError) as error:
        extract_blocks(mime, content)

    assert error.value.code == "file_too_large"
    assert limits.max_file_bytes_for(mime) == len(content) - 1
    assert limits.max_file_bytes_for("application/pdf") == limits.MAX_FILE_BYTES
