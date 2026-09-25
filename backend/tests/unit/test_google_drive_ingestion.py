from io import BytesIO
from threading import Event, Lock, Thread
from uuid import uuid4

import pytest
from cryptography.fernet import Fernet
from docx import Document as DocxDocument

from app.ingestion.google_drive import GoogleDriveDocumentProvider
from app.integrations.google_drive import (
    CredentialCipher,
    GoogleCredentials,
    GoogleDriveOAuthClient,
    GoogleRemoteUnauthorized,
    RemoteFile,
)
from app.workspaces.models import WorkspaceFolderSelection

GOOGLE_DOC = "application/vnd.google-apps.document"
PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class FakeGoogleDriveClient:
    def __init__(self, files: list[RemoteFile], content: dict[str, bytes], root_files: list[RemoteFile] | None = None, all_files: list[RemoteFile] | None = None) -> None:
        self.files = files
        self.content = content
        self.root_files = root_files if root_files is not None else files
        self.all_files = all_files if all_files is not None else files
        self.list_calls: list[str] = []
        self.read_calls: list[str] = []
        self.read_error: Exception | None = None

    def list_folder_files(self, *, credentials: GoogleCredentials, root_folder_id: str) -> list[RemoteFile]:
        self.list_calls.append(root_folder_id)
        return self.files

    def list_root_files(self, *, credentials: GoogleCredentials) -> list[RemoteFile]:
        self.list_calls.append("root")
        return self.root_files

    def list_all_files(self, *, credentials: GoogleCredentials) -> list[RemoteFile]:
        self.list_calls.append("all")
        return self.all_files

    def read_file(self, *, credentials: GoogleCredentials, remote_file: RemoteFile) -> bytes:
        self.read_calls.append(remote_file.id)
        if self.read_error is not None:
            raise self.read_error
        return self.content[remote_file.id]


def encrypted_credentials() -> tuple[CredentialCipher, str]:
    cipher = CredentialCipher(Fernet.generate_key().decode())
    return cipher, cipher.encrypt(GoogleCredentials("access-token", "refresh-token", None))


def remote_file(file_id: str, mime_type: str) -> RemoteFile:
    return RemoteFile(
        id=file_id,
        name=f"{file_id}.document",
        mime_type=mime_type,
        source_url=f"https://drive.example.test/{file_id}",
        modified_at=None,
    )


def selection(kind: str, external_folder_id: str = "") -> WorkspaceFolderSelection:
    return WorkspaceFolderSelection(id=uuid4(), workspace_folder_id=uuid4(), kind=kind, external_folder_id=external_folder_id)


def docx_bytes(text: str) -> bytes:
    document = DocxDocument()
    document.add_paragraph(text)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_provider_extracts_google_docs_and_docx_and_skips_unsupported_without_download() -> None:
    files = [remote_file("google-doc", GOOGLE_DOC), remote_file("word-doc", DOCX), remote_file("slides", "application/vnd.google-apps.presentation")]
    client = FakeGoogleDriveClient(
        files,
        {
            "google-doc": b"Approved marketing scope",
            "word-doc": docx_bytes("Approved consulting scope"),
        },
    )
    cipher, credentials = encrypted_credentials()

    discovered = GoogleDriveDocumentProvider(client, cipher).discover(
        encrypted_credentials=credentials,
        selections=[selection("folder", "root-folder")],
    )

    assert client.list_calls == ["root-folder"]
    assert set(client.read_calls) == {"google-doc", "word-doc"}
    assert [(document.external_file_id, document.text, document.error_code) for document in discovered] == [
        ("google-doc", "Approved marketing scope", None),
        ("slides", None, None),
        ("word-doc", "Approved consulting scope", None),
    ]


def test_provider_removes_nul_from_extracted_text_and_blocks() -> None:
    client = FakeGoogleDriveClient(
        [remote_file("google-doc", GOOGLE_DOC)],
        {"google-doc": b"Before\x00 after"},
    )
    cipher, credentials = encrypted_credentials()

    discovered = GoogleDriveDocumentProvider(client, cipher).discover(
        encrypted_credentials=credentials,
        selections=[selection("folder", "root-folder")],
    )

    assert discovered[0].text == "Before after"
    assert [block.text for block in discovered[0].blocks] == ["Before after"]


def test_provider_retains_safe_parent_metadata_for_library_projection() -> None:
    file = RemoteFile(
        id="brief",
        name="brief.document",
        mime_type=GOOGLE_DOC,
        source_url="https://drive.example.test/brief",
        modified_at=None,
        parent_ids=("campaign",),
    )
    client = FakeGoogleDriveClient([file], {"brief": b"Approved scope"})
    cipher, credentials = encrypted_credentials()

    discovered = GoogleDriveDocumentProvider(client, cipher).discover(
        encrypted_credentials=credentials,
        selections=[selection("folder", "root-folder")],
    )

    assert discovered[0].parent_ids == ("campaign",)


def test_structured_extraction_preserves_markdown_sections_and_docx_headings_tables() -> None:
    from app.ingestion.google_drive import _extract_blocks

    markdown = _extract_blocks("text/markdown", b"# Billing\n\nInvoices are due in 30 days.\n\n## Exceptions\n\nContact finance.")
    assert [block.section_path for block in markdown] == ["Billing", "Billing › Exceptions"]
    assert "Invoices are due" in markdown[0].text

    docx = DocxDocument()
    docx.add_heading("Policy", level=1)
    docx.add_paragraph("Employees submit expenses monthly.")
    table = docx.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Category"
    table.cell(0, 1).text = "Limit"
    table.cell(1, 0).text = "Travel"
    table.cell(1, 1).text = "$500"
    output = BytesIO()
    docx.save(output)
    docx_blocks = _extract_blocks(DOCX, output.getvalue())
    assert docx_blocks[0].section_path == "Policy"
    assert docx_blocks[1].section_path == "Policy"
    assert "Category: Travel" in docx_blocks[1].text and "Limit: $500" in docx_blocks[1].text


def test_provider_extracts_pdf_text(monkeypatch: pytest.MonkeyPatch) -> None:
    class Page:
        def extract_text(self) -> str:
            return "PDF source evidence"

    class FakePdfReader:
        def __init__(self) -> None:
            self.pages = [Page()]

    monkeypatch.setattr("app.ingestion.google_drive.PdfReader", lambda _: FakePdfReader())
    client = FakeGoogleDriveClient([remote_file("report", PDF)], {"report": b"not-real-pdf"})
    cipher, credentials = encrypted_credentials()

    discovered = GoogleDriveDocumentProvider(client, cipher).discover(
        encrypted_credentials=credentials,
        selections=[selection("folder", "root-folder")],
    )

    assert discovered[0].text == "PDF source evidence"
    assert discovered[0].error_code is None


def test_remote_file_authorization_error_becomes_a_document_failure() -> None:
    client = FakeGoogleDriveClient([remote_file("private-doc", GOOGLE_DOC)], {"private-doc": b"irrelevant"})
    client.read_error = GoogleRemoteUnauthorized()
    cipher, credentials = encrypted_credentials()

    discovered = GoogleDriveDocumentProvider(client, cipher).discover(
        encrypted_credentials=credentials,
        selections=[selection("folder", "root-folder")],
    )

    assert discovered[0].error_code == "source_file_unavailable"


def test_provider_unions_overlapping_folder_selections_before_downloading() -> None:
    item = remote_file("shared-doc", GOOGLE_DOC)
    client = FakeGoogleDriveClient([item], {"shared-doc": b"Only once"})
    cipher, credentials = encrypted_credentials()

    discovered = GoogleDriveDocumentProvider(client, cipher).discover(
        encrypted_credentials=credentials,
        selections=[selection("folder", "parent"), selection("folder", "child")],
    )

    assert client.list_calls == ["parent", "child"]
    assert client.read_calls == ["shared-doc"]
    assert [item.external_file_id for item in discovered] == ["shared-doc"]


def test_provider_combines_direct_root_files_without_traversing_root_folders() -> None:
    folder_file = remote_file("folder-doc", GOOGLE_DOC)
    root_file = remote_file("root-doc", GOOGLE_DOC)
    client = FakeGoogleDriveClient(
        [folder_file],
        {"folder-doc": b"In selected folder", "root-doc": b"At root"},
        root_files=[root_file],
    )
    cipher, credentials = encrypted_credentials()

    discovered = GoogleDriveDocumentProvider(client, cipher).discover(
        encrypted_credentials=credentials,
        selections=[selection("folder", "project"), selection("root_files")],
    )

    assert client.list_calls == ["project", "root"]
    assert {item.external_file_id for item in discovered} == {"folder-doc", "root-doc"}


def test_provider_all_accessible_uses_one_complete_listing() -> None:
    item = remote_file("anywhere", GOOGLE_DOC)
    client = FakeGoogleDriveClient([], {"anywhere": b"Accessible"}, all_files=[item])
    cipher, credentials = encrypted_credentials()

    discovered = GoogleDriveDocumentProvider(client, cipher).discover(
        encrypted_credentials=credentials,
        selections=[selection("all_accessible")],
    )

    assert client.list_calls == ["all"]
    assert [item.external_file_id for item in discovered] == ["anywhere"]


def test_provider_extracts_with_bounded_concurrency_and_stable_order() -> None:
    files = [remote_file(f"file-{index}", GOOGLE_DOC) for index in reversed(range(7))]

    class BlockingClient(FakeGoogleDriveClient):
        def __init__(self) -> None:
            super().__init__(files, {item.id: item.id.encode() for item in files})
            self.lock = Lock()
            self.active_reads = 0
            self.max_active_reads = 0
            self.six_reads_started = Event()
            self.release_reads = Event()

        def read_file(self, *, credentials: GoogleCredentials, remote_file: RemoteFile) -> bytes:
            with self.lock:
                self.active_reads += 1
                self.max_active_reads = max(self.max_active_reads, self.active_reads)
                if self.active_reads == 6:
                    self.six_reads_started.set()
            self.release_reads.wait(timeout=2)
            try:
                return super().read_file(credentials=credentials, remote_file=remote_file)
            finally:
                with self.lock:
                    self.active_reads -= 1

    client = BlockingClient()
    cipher, credentials = encrypted_credentials()
    discovered = []
    thread = Thread(
        target=lambda: discovered.extend(
            GoogleDriveDocumentProvider(client, cipher, extraction_workers=100).discover(
                encrypted_credentials=credentials,
                selections=[selection("all_accessible")],
            )
        )
    )

    thread.start()
    assert client.six_reads_started.wait(timeout=2)
    client.release_reads.set()
    thread.join(timeout=2)

    assert not thread.is_alive()
    assert client.max_active_reads == 6
    assert [item.external_file_id for item in discovered] == sorted(item.id for item in files)
    assert [item.text for item in discovered] == sorted(item.id for item in files)


def test_concurrent_extraction_keeps_per_file_failures_without_cancelling_siblings() -> None:
    files = [
        remote_file("good-a", GOOGLE_DOC),
        remote_file("private", GOOGLE_DOC),
        remote_file("bad", GOOGLE_DOC),
        remote_file("slides", "application/vnd.google-apps.presentation"),
        remote_file("good-b", GOOGLE_DOC),
    ]

    class SelectivelyFailingClient(FakeGoogleDriveClient):
        def read_file(self, *, credentials: GoogleCredentials, remote_file: RemoteFile) -> bytes:
            if remote_file.id == "private":
                raise GoogleRemoteUnauthorized()
            if remote_file.id == "bad":
                raise ValueError("invalid source")
            return super().read_file(credentials=credentials, remote_file=remote_file)

    client = SelectivelyFailingClient(
        files,
        {"good-a": b"Good A", "good-b": b"Good B", "private": b"", "bad": b""},
    )
    cipher, credentials = encrypted_credentials()

    discovered = GoogleDriveDocumentProvider(client, cipher, extraction_workers=2).discover(
        encrypted_credentials=credentials,
        selections=[selection("all_accessible")],
    )

    assert [(item.external_file_id, item.text, item.error_code) for item in discovered] == [
        ("bad", None, "text_extraction_failed"),
        ("good-a", "Good A", None),
        ("good-b", "Good B", None),
        ("private", None, "source_file_unavailable"),
        ("slides", None, None),
    ]
    assert "slides" not in client.read_calls


def test_google_client_lists_direct_root_files_with_pagination(monkeypatch: pytest.MonkeyPatch) -> None:
    calls: list[dict[str, object]] = []

    class Response:
        status_code = 200

        def __init__(self, body: dict[str, object]) -> None:
            self.body = body

        def raise_for_status(self) -> None:
            return None

        def json(self) -> dict[str, object]:
            return self.body

    pages = [
        {"nextPageToken": "next", "files": [{"id": "root-a", "name": "A", "mimeType": GOOGLE_DOC}]},
        {"files": [{"id": "root-b", "name": "B", "mimeType": PDF}]},
    ]

    def fake_get(_: str, *, params: dict[str, object], **__: object) -> Response:
        calls.append(params)
        return Response(pages.pop(0))

    monkeypatch.setattr("app.integrations.google_drive.httpx.get", fake_get)
    files = GoogleDriveOAuthClient(client_id="id", client_secret="secret", redirect_uri="https://example.test/callback").list_root_files(
        credentials=GoogleCredentials("access", None, None)
    )

    assert [item.id for item in files] == ["root-a", "root-b"]
    assert calls[0]["q"] == "'root' in parents and trashed = false and mimeType != 'application/vnd.google-apps.folder'"
    assert calls[0]["pageToken"] is None
    assert calls[1]["pageToken"] == "next"
    assert calls[0]["supportsAllDrives"] == "true"
