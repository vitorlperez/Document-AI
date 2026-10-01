"""Extractors take a seekable file-like (a TemporaryFile from the ranged download), not only bytes."""

import tempfile
from contextlib import contextmanager
from io import BytesIO

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook
from test_ocr_pdf_layers import FakeOcr, pdf

from app.ingestion.extraction import extract_blocks, limits
from app.ingestion.extraction.errors import ExtractionError
from app.ingestion.extraction.mime import DOCX, PDF, XLSX


@contextmanager
def spooled(content: bytes):
    with tempfile.TemporaryFile() as handle:
        handle.write(content)
        handle.seek(5)  # extractors must not depend on the caller's position
        yield handle


def docx_bytes() -> bytes:
    document, out = DocxDocument(), BytesIO()
    document.add_paragraph("Cláusula de rescisão")
    document.save(out)
    return out.getvalue()


def xlsx_bytes() -> bytes:
    book, out = Workbook(), BytesIO()
    book.active.append(["nome", "valor"])
    book.active.append(["Acme", 10])
    book.save(out)
    return out.getvalue()


@pytest.mark.parametrize("mime,content,needle", [
    (DOCX, docx_bytes(), "rescisão"),
    (XLSX, xlsx_bytes(), "Acme"),
    (PDF, pdf(["Contrato de prestacao de servicos numero 42"]), "Contrato"),
    ("text/plain", b"texto simples\n\noutro", "texto simples"),
])
def test_file_like_and_bytes_give_the_same_blocks(mime, content, needle):
    with spooled(content) as handle:
        from_stream = extract_blocks(mime, handle)
    assert from_stream == extract_blocks(mime, content)
    assert needle in " ".join(b.text for b in from_stream)


def test_scanned_pdf_ocr_from_a_file_like_uses_batches():
    ocr = FakeOcr()
    with spooled(pdf([None] * 10)) as handle:
        blocks = extract_blocks(PDF, handle, ocr=ocr)
    assert ocr.calls == [8, 2] and all(b.ocr for b in blocks)


def test_size_guard_reads_the_file_like_size_not_its_bytes(monkeypatch):
    monkeypatch.setattr(limits, "MAX_FILE_BYTES", 10)
    with spooled(docx_bytes()) as handle, pytest.raises(ExtractionError) as caught:
        extract_blocks(DOCX, handle)
    assert caught.value.code == "file_too_large"


def test_default_limit_is_100_mib():
    assert limits.MAX_FILE_BYTES == 100 * 1024 * 1024
