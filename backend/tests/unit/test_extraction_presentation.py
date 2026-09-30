from io import BytesIO

from pptx import Presentation
from pptx.util import Inches

from app.ingestion.extraction import extract_blocks

PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def deck() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    slide.shapes.title.text = "Proposta comercial"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
    box.text_frame.text = "Prazo de 30 dias"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(5), Inches(1)).table
    for (r, c), v in {(0, 0): "Item", (0, 1): "Preço", (1, 0): "Licença", (1, 1): "R$ 100"}.items():
        table.cell(r, c).text = v
    slide.notes_slide.notes_text_frame.text = "Lembrar do desconto"
    out = BytesIO()
    prs.save(out)
    return out.getvalue()


def test_slide_text_table_and_notes_are_extracted_with_slide_numbers() -> None:
    blocks = extract_blocks(PPTX, deck())
    body = next(b for b in blocks if b.section_path == "Proposta comercial")
    assert body.page_number == 1
    assert "Prazo de 30 dias" in body.text and "Item: Licença | Preço: R$ 100" in body.text
    notes = next(b for b in blocks if b.section_path == "Proposta comercial › Notas do orador")
    assert notes.page_number == 1 and notes.text == "Lembrar do desconto"
