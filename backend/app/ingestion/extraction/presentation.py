from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.ingestion.blocks import ExtractedBlock
from app.ingestion.extraction import limits


def _shapes(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shapes(shape.shapes)
        else:
            yield shape


def pptx_blocks(content) -> list[ExtractedBlock]:
    limits.guard_size(content)
    limits.guard_zip(content)
    presentation = Presentation(limits.as_stream(content))
    blocks: list[ExtractedBlock] = []
    for number, slide in enumerate(presentation.slides, start=1):
        title_shape = (
            slide.shapes.title
        )  # a new proxy on every access: compare by shape_id, not identity
        title = title_shape.text.strip() if title_shape is not None else ""
        title = title or f"Slide {number}"
        title_id = title_shape.shape_id if title_shape is not None else None
        lines: list[str] = []
        for shape in _shapes(slide.shapes):
            if shape.has_text_frame and shape.shape_id != title_id:
                lines.extend(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
            elif getattr(shape, "has_table", False) and shape.has_table:
                rows = [
                    [" ".join(cell.text.split()) for cell in row.cells] for row in shape.table.rows
                ]
                header = rows[0] if rows else []
                for row in rows[1:] or rows:
                    lines.append(
                        " | ".join(
                            f"{header[i]}: {v}" if i < len(header) and header[i] and rows[1:] else v
                            for i, v in enumerate(row)
                            if v
                        )
                    )
        if lines:
            blocks.append(ExtractedBlock("\n".join(lines), page_number=number, section_path=title))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            blocks.append(
                ExtractedBlock(
                    slide.notes_slide.notes_text_frame.text.strip(),
                    page_number=number,
                    section_path=f"{title} › Notas do orador",
                )
            )
    return blocks
