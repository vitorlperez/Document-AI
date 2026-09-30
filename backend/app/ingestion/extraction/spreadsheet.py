"""CSV/XLSX -> header-carrying row blocks (citable as 'Aba › linhas a–b')."""

import csv
import io
from collections.abc import Iterable, Sequence
from datetime import date, datetime

from openpyxl import load_workbook

from app.ingestion.blocks import ExtractedBlock
from app.ingestion.extraction import limits
from app.ingestion.extraction.text import decode_text


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        stamp = value.strftime("%d/%m/%Y")
        return (
            stamp
            if (value.hour, value.minute, value.second) == (0, 0, 0)
            else f"{stamp} {value:%H:%M}"
        )
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else f"{value:.10g}"
    return " ".join(str(value).split())[: limits.MAX_CELL_CHARS]


def rows_to_blocks(title: str, rows: Iterable[Sequence[object]]) -> list[ExtractedBlock]:
    header: list[str] | None = None
    blocks: list[ExtractedBlock] = []
    lines: list[str] = []
    first = last = 0
    truncated = False

    def flush() -> None:
        nonlocal lines
        if lines:
            blocks.append(
                ExtractedBlock("\n".join(lines), section_path=f"{title} › linhas {first}–{last}")
            )
            lines = []

    for number, raw in enumerate(rows, start=1):
        if number > limits.MAX_ROWS_PER_SHEET + 1:
            truncated = True
            break
        cells = [_cell(v) for v in list(raw)[: limits.MAX_COLUMNS]]
        if not any(cells):
            continue
        if header is None:
            header = cells
            continue
        line = " | ".join(
            f"{header[i]}: {v}" if i < len(header) and header[i] else v
            for i, v in enumerate(cells)
            if v
        )
        if not lines:
            first = number
        lines.append(line)
        last = number
        if len(lines) == limits.ROWS_PER_BLOCK:
            flush()
    flush()
    if header is not None and not blocks:  # a sheet with only one row: keep it searchable
        blocks.append(ExtractedBlock(" | ".join(v for v in header if v), section_path=title))
    if truncated:
        blocks.append(
            ExtractedBlock(
                f"[Conteúdo truncado: apenas as primeiras {limits.MAX_ROWS_PER_SHEET} linhas da aba foram indexadas.]",
                section_path=title,
            )
        )
    return blocks


def csv_blocks(content: bytes) -> list[ExtractedBlock]:
    limits.guard_size(content)
    text = decode_text(content)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if sample.count("\t") > sample.count(",") else csv.excel
        if sample.count(";") > sample.count(","):
            dialect = type("Semicolon", (csv.excel,), {"delimiter": ";"})
    return rows_to_blocks("Tabela", csv.reader(io.StringIO(text), dialect))


def xlsx_blocks(content: bytes) -> list[ExtractedBlock]:
    limits.guard_size(content)
    limits.guard_zip(content)
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        blocks: list[ExtractedBlock] = []
        for sheet in workbook.worksheets[: limits.MAX_SHEETS]:
            if sheet.sheet_state != "visible":
                continue
            sheet.reset_dimensions()  # read-only mode trusts a possibly wrong <dimension> tag
            blocks.extend(rows_to_blocks(sheet.title, sheet.iter_rows(values_only=True)))
        return blocks
    finally:
        workbook.close()
