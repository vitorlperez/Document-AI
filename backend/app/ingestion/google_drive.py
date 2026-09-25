"""Google Drive discovery and text extraction behind the ingestion boundary."""

import re
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from zipfile import BadZipFile

from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.ingestion.service import ELIGIBLE_MIME_TYPES, DiscoveredDocument, ExtractedBlock
from app.integrations.google_drive import (
    CredentialCipher,
    GoogleCredentials,
    GoogleDriveOAuthClient,
    GoogleRemoteUnauthorized,
    RemoteFile,
    RemoteFolder,
)
from app.workspaces.models import WorkspaceFolderSelection

MAX_EXTRACTION_WORKERS = 6


class GoogleDriveDocumentProvider:
    """Fetch an authorized scope union, then retain extracted text only in the DB."""

    def __init__(
        self,
        client: GoogleDriveOAuthClient,
        cipher: CredentialCipher,
        *,
        extraction_workers: int = MAX_EXTRACTION_WORKERS,
    ):
        self.client = client
        self.cipher = cipher
        self.extraction_workers = min(MAX_EXTRACTION_WORKERS, max(1, extraction_workers))

    def discover(
        self,
        *,
        encrypted_credentials: str,
        selections: list[WorkspaceFolderSelection],
    ) -> list[DiscoveredDocument]:
        credentials = self.cipher.decrypt(encrypted_credentials)
        remote_files: dict[str, RemoteFile] = {}
        for selection in selections:
            if selection.kind == "folder":
                found = self.client.list_folder_files(
                    credentials=credentials,
                    root_folder_id=selection.external_folder_id,
                )
            elif selection.kind == "root_files":
                found = self.client.list_root_files(credentials=credentials)
            elif selection.kind == "all_accessible":
                found = self.client.list_all_files(credentials=credentials)
            else:
                raise ValueError("unsupported workspace selection")
            for remote_file in found:
                remote_files.setdefault(remote_file.id, remote_file)
        # A Drive response does not promise an implicit ordering. Stable IDs
        # make repeated full snapshots and the concurrent projection predictable.
        files = sorted(remote_files.values(), key=lambda remote_file: remote_file.id)
        if not files:
            return []
        # Only remote reads and parsing happen on threads. The caller keeps all
        # SQLAlchemy access in the Celery worker thread after this returns.
        with ThreadPoolExecutor(max_workers=min(self.extraction_workers, len(files))) as executor:
            return list(
                executor.map(
                    lambda remote_file: self._extract(credentials=credentials, remote_file=remote_file),
                    files,
                )
            )

    def folders(self, *, encrypted_credentials: str) -> list[RemoteFolder]:
        """Return safe folder metadata for the Company Library projection.

        This is intentionally separate from extraction: the library retains no
        provider credentials and does not browse Drive during a user request.
        """
        return self.client.list_folders(credentials=self.cipher.decrypt(encrypted_credentials))

    def _extract(self, *, credentials: GoogleCredentials, remote_file: RemoteFile) -> DiscoveredDocument:
        base = {
            "external_file_id": remote_file.id,
            "name": remote_file.name,
            "mime_type": remote_file.mime_type,
            "source_url": remote_file.source_url,
            "modified_at": remote_file.modified_at,
            "parent_ids": remote_file.parent_ids,
        }
        if remote_file.mime_type not in ELIGIBLE_MIME_TYPES:
            return DiscoveredDocument(**base)
        try:
            content = self.client.read_file(credentials=credentials, remote_file=remote_file)
            blocks = _extract_blocks(remote_file.mime_type, content)
        except GoogleRemoteUnauthorized:
            # A listing token can remain valid while one shared/export-restricted
            # file rejects its content request. Treat that as an item failure;
            # source-level authorization failures are still raised by discovery.
            return DiscoveredDocument(**base, error_code="source_file_unavailable")
        except (BadZipFile, PackageNotFoundError, PdfReadError, UnicodeDecodeError, ValueError):
            # Do not attach the remote body or exception to records/logs. The
            # document's explicit code is enough for an Admin to retry safely.
            return DiscoveredDocument(**base, error_code="text_extraction_failed")
        # Google exports and PDF extractors can return NUL characters. PostgreSQL
        # rejects them in text fields, which would otherwise fail the whole sync.
        blocks = [
            ExtractedBlock(
                block.text.replace("\x00", ""),
                page_number=block.page_number,
                section_path=block.section_path.replace("\x00", "") if block.section_path else None,
            )
            for block in blocks
        ]
        return DiscoveredDocument(**base, text="\n\n".join(block.text for block in blocks), blocks=tuple(blocks))


def _extract_blocks(mime_type: str, content: bytes) -> list[ExtractedBlock]:
    if mime_type == "application/pdf":
        return [
            ExtractedBlock(page.extract_text() or "", page_number=index)
            for index, page in enumerate(PdfReader(BytesIO(content)).pages, start=1)
        ]
    if mime_type == "application/vnd.google-apps.document":
        return [ExtractedBlock(paragraph) for paragraph in content.decode("utf-8").splitlines()]
    if mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        docx = DocxDocument(BytesIO(content))
        blocks: list[ExtractedBlock] = []
        headings: list[tuple[int, str]] = []
        for element in docx.element.body.iterchildren():
            if isinstance(element, CT_P):
                paragraph = Paragraph(element, docx)
                if paragraph.style and paragraph.style.name.lower().startswith("heading") and paragraph.text.strip():
                    level_match = re.search(r"(\d+)$", paragraph.style.name)
                    level = int(level_match.group(1)) if level_match else 1
                    headings = [(prior, title) for prior, title in headings if prior < level]
                    headings.append((level, paragraph.text.strip()))
                elif paragraph.text.strip():
                    blocks.append(ExtractedBlock(paragraph.text, section_path=" › ".join(title for _, title in headings) or None))
            elif isinstance(element, CT_Tbl):
                table = Table(element, docx)
                headers = [" ".join(cell.text.split()) for cell in table.rows[0].cells] if table.rows else []
                for row in (table.rows[1:] if len(table.rows) > 1 else table.rows):
                    cells = [" ".join(cell.text.split()) for cell in row.cells]
                    if any(cells):
                        row_text = " | ".join(
                            f"{headers[index]}: {value}" if index < len(headers) and headers[index] else value
                            for index, value in enumerate(cells)
                        )
                        blocks.append(ExtractedBlock(row_text, section_path=" › ".join(title for _, title in headings) or None))
        return blocks
    if mime_type == "text/markdown":
        return _markdown_blocks(content.decode("utf-8"))
    raise ValueError("unsupported file type")


def _markdown_blocks(text: str) -> list[ExtractedBlock]:
    blocks: list[ExtractedBlock] = []
    headings: list[tuple[int, str]] = []
    current: list[str] = []
    for line in text.splitlines():
        match = re.match(r"^(#{1,6})\s+(.+)$", line.strip())
        if match:
            if current:
                blocks.append(ExtractedBlock("\n".join(current), section_path=" › ".join(name for _, name in headings)))
                current = []
            level, title = len(match.group(1)), match.group(2).strip()
            headings = [(prior_level, name) for prior_level, name in headings if prior_level < level]
            headings.append((level, title))
        elif line.strip():
            current.append(line)
        elif current:
            blocks.append(ExtractedBlock("\n".join(current), section_path=" › ".join(name for _, name in headings) or None))
            current = []
    if current:
        blocks.append(ExtractedBlock("\n".join(current), section_path=" › ".join(name for _, name in headings) or None))
    return blocks


def _extract_text(mime_type: str, content: bytes) -> str:
    return "\n\n".join(block.text for block in _extract_blocks(mime_type, content))
