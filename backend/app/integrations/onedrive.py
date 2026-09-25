"""Delegated Microsoft identity and OneDrive Graph adapter."""

import json
import logging
import secrets
import time
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from datetime import UTC, datetime, timedelta
from typing import Any
from urllib.parse import quote, urlencode, urlparse
from uuid import UUID
from zipfile import BadZipFile

import httpx
from cryptography.fernet import Fernet, InvalidToken
from docx.opc.exceptions import PackageNotFoundError
from pypdf.errors import PdfReadError
from sqlalchemy import select, update
from sqlalchemy.orm import Session

from app.audit_usage.models import AuditLog
from app.core.scoping import OrganizationScope
from app.identity.auth import hash_secret
from app.identity.models import UserSession
from app.ingestion.service import ELIGIBLE_MIME_TYPES, DiscoveredDocument, DiscoveryResult
from app.integrations.errors import SourceRemoteUnauthorized
from app.integrations.google_drive import RemoteFolder
from app.integrations.models import DataSource, OAuthConnectionState
from app.organizations.models import Membership, MembershipRole
from app.workspaces.models import WorkspaceFolder, WorkspaceFolderSelection

logger = logging.getLogger("document_intelligence.integration")
# The common audience accepts both Microsoft personal and work/school accounts.
MICROSOFT_AUTHORITY = "https://login.microsoftonline.com/common/oauth2/v2.0"
GRAPH_ROOT = "https://graph.microsoft.com/v1.0"
GRAPH_SCOPES = "openid profile offline_access User.Read Files.Read"
MAX_PAGE_WORKERS = 4


class OneDriveOAuthUnavailable(RuntimeError):
    pass


class OneDriveOAuthInvalid(ValueError):
    pass


class OneDriveAccessDenied(PermissionError):
    pass


class OneDriveReauthRequired(OneDriveOAuthInvalid):
    pass


class OneDriveAccountMismatch(ValueError):
    def __init__(self, organization_id: UUID):
        self.organization_id = organization_id
        super().__init__("reconnect with the existing OneDrive account")


class OneDriveCursorInvalid(OneDriveOAuthInvalid):
    pass


class OneDriveDeltaExpired(RuntimeError):
    pass


@dataclass(frozen=True)
class OneDriveCredentials:
    access_token: str
    refresh_token: str
    expires_at: datetime


@dataclass(frozen=True)
class DeltaPage:
    items: list[dict[str, Any]]
    delta_link: str


class OneDriveCipher:
    def __init__(self, key: str | None):
        self.key = key

    def _fernet(self) -> Fernet:
        if not self.key:
            raise OneDriveOAuthUnavailable("OneDrive encryption is not configured")
        try:
            return Fernet(self.key.encode())
        except (ValueError, TypeError) as error:
            raise OneDriveOAuthUnavailable("OneDrive encryption is not configured") from error

    def encrypt_credentials(self, credentials: OneDriveCredentials) -> str:
        payload = json.dumps(
            {
                "access_token": credentials.access_token,
                "refresh_token": credentials.refresh_token,
                "expires_at": credentials.expires_at.isoformat(),
            },
            separators=(",", ":"),
        )
        return self._fernet().encrypt(payload.encode()).decode()

    def decrypt_credentials(self, value: str) -> OneDriveCredentials:
        try:
            payload = json.loads(self._fernet().decrypt(value.encode()))
            return OneDriveCredentials(
                access_token=str(payload["access_token"]),
                refresh_token=str(payload["refresh_token"]),
                expires_at=datetime.fromisoformat(str(payload["expires_at"])),
            )
        except (InvalidToken, KeyError, TypeError, ValueError, json.JSONDecodeError) as error:
            raise OneDriveOAuthInvalid("OneDrive credentials are unavailable") from error

    def encrypt_cursor(self, value: str) -> str:
        return self._fernet().encrypt(value.encode()).decode()

    def decrypt_cursor(self, value: str) -> str:
        try:
            return self._fernet().decrypt(value.encode()).decode()
        except (InvalidToken, ValueError) as error:
            raise OneDriveCursorInvalid("OneDrive sync cursor is unavailable") from error


class MicrosoftGraphClient:
    def __init__(
        self, *, client_id: str | None, client_secret: str | None, redirect_uri: str | None
    ):
        self.client_id, self.client_secret, self.redirect_uri = (
            client_id,
            client_secret,
            redirect_uri,
        )

    def _configured(self) -> None:
        if not self.client_id or not self.client_secret or not self.redirect_uri:
            raise OneDriveOAuthUnavailable("Microsoft OAuth is not configured")

    def authorization_url(self, *, state: str) -> str:
        self._configured()
        return f"{MICROSOFT_AUTHORITY}/authorize?" + urlencode(
            {
                "client_id": self.client_id,
                "response_type": "code",
                "redirect_uri": self.redirect_uri,
                "response_mode": "query",
                "prompt": "select_account",
                "scope": GRAPH_SCOPES,
                "state": state,
            }
        )

    def exchange_code(self, *, code: str) -> OneDriveCredentials:
        self._configured()
        return self._token_request(
            {
                "grant_type": "authorization_code",
                "code": code,
                "redirect_uri": self.redirect_uri,
                "scope": GRAPH_SCOPES,
            }
        )

    def refresh(self, *, credentials: OneDriveCredentials) -> OneDriveCredentials:
        self._configured()
        return self._token_request(
            {
                "grant_type": "refresh_token",
                "refresh_token": credentials.refresh_token,
                "scope": GRAPH_SCOPES,
            },
            previous=credentials,
        )

    def _token_request(
        self, data: dict[str, str | None], *, previous: OneDriveCredentials | None = None
    ) -> OneDriveCredentials:
        response = httpx.post(
            f"{MICROSOFT_AUTHORITY}/token",
            data={**data, "client_id": self.client_id, "client_secret": self.client_secret},
            timeout=15,
        )
        if response.status_code in {400, 401, 403}:
            raise OneDriveOAuthInvalid("Microsoft authorization failed")
        response.raise_for_status()
        payload = response.json()
        try:
            return OneDriveCredentials(
                access_token=str(payload["access_token"]),
                refresh_token=str(
                    payload.get("refresh_token") or (previous.refresh_token if previous else "")
                ),
                expires_at=datetime.now(UTC) + timedelta(seconds=int(payload["expires_in"])),
            )
        except (KeyError, TypeError, ValueError) as error:
            raise OneDriveOAuthInvalid("Microsoft authorization response is invalid") from error

    def account_email(self, *, credentials: OneDriveCredentials) -> str | None:
        data = self._get_json(
            f"{GRAPH_ROOT}/me?$select=mail,userPrincipalName", credentials=credentials
        )
        value = data.get("mail") or data.get("userPrincipalName")
        return str(value).strip().lower() if isinstance(value, str) and value.strip() else None

    def drive_id(self, *, credentials: OneDriveCredentials) -> str:
        data = self._get_json(f"{GRAPH_ROOT}/me/drive?$select=id", credentials=credentials)
        drive_id = data.get("id")
        if not isinstance(drive_id, str) or not drive_id:
            raise OneDriveOAuthInvalid("Microsoft drive response is invalid")
        return drive_id

    def root_id(self, *, credentials: OneDriveCredentials) -> str:
        data = self._get_json(f"{GRAPH_ROOT}/me/drive/root?$select=id", credentials=credentials)
        root_id = data.get("id")
        if not isinstance(root_id, str) or not root_id:
            raise OneDriveOAuthInvalid("Microsoft root response is invalid")
        return root_id

    def folders(self, *, credentials: OneDriveCredentials) -> list[RemoteFolder]:
        root_id = self.root_id(credentials=credentials)
        pending = [root_id]
        visited: set[str] = set()
        folders: list[RemoteFolder] = []
        while pending:
            parent_id = pending.pop()
            if parent_id in visited:
                continue
            visited.add(parent_id)
            for item in self._all_pages(
                f"{GRAPH_ROOT}/me/drive/items/{quote(parent_id, safe='')}/children?$select=id,name,folder,parentReference",
                credentials=credentials,
            ):
                item_id = str(item.get("id") or "")
                facet = item.get("folder")
                if not item_id or not isinstance(facet, dict):
                    continue
                composite_id = self._external_id(item, fallback_drive_id=None)
                parent_reference = item.get("parentReference") or {}
                parents = self._parent_ids(parent_reference)
                folders.append(
                    RemoteFolder(composite_id, str(item.get("name") or "Untitled"), parents)
                )
                pending.append(item_id)
        return folders

    def delta(
        self,
        *,
        credentials: OneDriveCredentials,
        selection: WorkspaceFolderSelection,
        cursor: str | None,
    ) -> DeltaPage:
        if cursor:
            url = cursor
        elif selection.kind == "folder":
            folder_id = quote(selection.external_folder_id, safe="")
            url = f"{GRAPH_ROOT}/me/drive/items/{folder_id}/delta?$select=id,name,file,folder,parentReference,webUrl,lastModifiedDateTime,deleted"
        else:
            url = f"{GRAPH_ROOT}/me/drive/root/delta?$select=id,name,file,folder,parentReference,webUrl,lastModifiedDateTime,deleted"
        items: list[dict[str, Any]] = []
        while url:
            data = self._get_json(url, credentials=credentials, allow_expired_delta=True)
            page_items = data.get("value") or []
            if isinstance(page_items, list):
                items.extend(item for item in page_items if isinstance(item, dict))
            next_link = data.get("@odata.nextLink")
            delta_link = data.get("@odata.deltaLink")
            if isinstance(next_link, str):
                url = self._validate_graph_url(next_link)
                continue
            if isinstance(delta_link, str):
                return DeltaPage(items=items, delta_link=self._validate_graph_url(delta_link))
            raise OneDriveOAuthInvalid("Microsoft delta response is invalid")
        raise OneDriveOAuthInvalid("Microsoft delta response is incomplete")

    def list_files(
        self, *, credentials: OneDriveCredentials, selection: WorkspaceFolderSelection
    ) -> list[dict[str, Any]]:
        if selection.kind == "folder":
            pending = [selection.external_folder_id]
            endpoint = lambda parent: (
                f"{GRAPH_ROOT}/me/drive/items/{quote(parent, safe='')}/children?$select=id,name,file,folder,parentReference,webUrl,lastModifiedDateTime"
            )
        elif selection.kind == "root_files" or selection.kind == "all_accessible":
            pending = [self.root_id(credentials=credentials)]
            endpoint = lambda parent: (
                f"{GRAPH_ROOT}/me/drive/items/{quote(parent, safe='')}/children?$select=id,name,file,folder,parentReference,webUrl,lastModifiedDateTime"
            )
        else:
            raise ValueError("unsupported workspace selection")
        files: dict[str, dict[str, Any]] = {}
        visited: set[str] = set()
        while pending:
            parent_id = pending.pop()
            if parent_id in visited:
                continue
            visited.add(parent_id)
            try:
                page_items = self._all_pages(endpoint(parent_id), credentials=credentials)
            except httpx.HTTPStatusError as error:
                if selection.kind == "folder" and error.response.status_code == 404:
                    continue
                raise
            for item in page_items:
                if isinstance(item.get("folder"), dict):
                    if selection.kind != "root_files":
                        pending.append(str(item.get("id") or ""))
                    continue
                key = self._external_id(item, fallback_drive_id=None)
                files[key] = item
        return list(files.values())

    def read_file(self, *, credentials: OneDriveCredentials, item_id: str) -> bytes:
        url = f"{GRAPH_ROOT}/me/drive/items/{quote(item_id, safe='')}/content"
        for attempt in range(4):
            response = httpx.get(
                url,
                headers={"Authorization": f"Bearer {credentials.access_token}"},
                timeout=30,
                follow_redirects=True,
            )
            if response.status_code in {401, 403}:
                raise SourceRemoteUnauthorized()
            if response.status_code == 429 and attempt < 3:
                retry_after = response.headers.get("Retry-After", "1")
                try:
                    time.sleep(min(max(float(retry_after), 0.0), 30.0))
                except ValueError:
                    time.sleep(1)
                continue
            response.raise_for_status()
            return response.content
        raise RuntimeError("Microsoft Graph throttling limit reached")

    def _all_pages(
        self, url: str, *, credentials: OneDriveCredentials | None = None
    ) -> list[dict[str, Any]]:
        if credentials is None:
            raise ValueError("credentials are required")
        rows: list[dict[str, Any]] = []
        while url:
            data = self._get_json(url, credentials=credentials)
            values = data.get("value") or []
            if isinstance(values, list):
                rows.extend(item for item in values if isinstance(item, dict))
            next_link = data.get("@odata.nextLink")
            url = self._validate_graph_url(next_link) if isinstance(next_link, str) else ""
        return rows

    @staticmethod
    def _validate_graph_url(url: str) -> str:
        parsed = urlparse(url)
        if (
            parsed.scheme != "https"
            or parsed.netloc.lower() != "graph.microsoft.com"
            or not parsed.path.startswith("/v1.0/")
        ):
            raise OneDriveOAuthInvalid("Microsoft returned an invalid pagination link")
        return url

    def _get_json(
        self, url: str, *, credentials: OneDriveCredentials, allow_expired_delta: bool = False
    ) -> dict[str, Any]:
        self._validate_graph_url(url)
        for attempt in range(4):
            response = httpx.get(
                url, headers={"Authorization": f"Bearer {credentials.access_token}"}, timeout=20
            )
            if response.status_code in {401, 403}:
                raise SourceRemoteUnauthorized()
            if allow_expired_delta and response.status_code == 410:
                raise OneDriveDeltaExpired()
            if allow_expired_delta and response.status_code == 404:
                try:
                    error_payload = response.json().get("error", {})
                except (ValueError, AttributeError):
                    error_payload = {}
                if isinstance(error_payload, dict) and error_payload.get("code") == "itemNotFound":
                    raise OneDriveDeltaExpired()
            if response.status_code == 429 and attempt < 3:
                retry_after = response.headers.get("Retry-After", "1")
                try:
                    time.sleep(min(max(float(retry_after), 0.0), 30.0))
                except ValueError:
                    time.sleep(1)
                continue
            response.raise_for_status()
            data = response.json()
            return data if isinstance(data, dict) else {}
        raise RuntimeError("Microsoft Graph throttling limit reached")

    @staticmethod
    def _external_id(item: dict[str, Any], *, fallback_drive_id: str | None = None) -> str:
        item_id = str(item.get("id") or "")
        if not item_id:
            raise ValueError("Microsoft item metadata is incomplete")
        # The source is bound to one user's default drive in this release, so
        # the Graph item ID is unique within this source's namespace.
        return item_id

    @staticmethod
    def _parent_ids(parent: dict[str, Any]) -> tuple[str, ...]:
        item_id = parent.get("id")
        path = str(parent.get("path") or "").rstrip("/")
        if item_id == "root" or path == "/drive/root:":
            return ("root",)
        if item_id:
            return (str(item_id),)
        return ()


class OneDriveDocumentProvider:
    key = "onedrive"

    def __init__(self, client: MicrosoftGraphClient, cipher: OneDriveCipher):
        self.client, self.cipher = client, cipher
        self.updated_encrypted_credentials: str | None = None

    def _credentials(self, encrypted_credentials: str | None) -> OneDriveCredentials:
        try:
            credentials = self.cipher.decrypt_credentials(
                self.updated_encrypted_credentials or encrypted_credentials or ""
            )
        except OneDriveOAuthInvalid as error:
            raise SourceRemoteUnauthorized() from error
        if credentials.expires_at <= datetime.now(UTC) + timedelta(minutes=5):
            try:
                credentials = self.client.refresh(credentials=credentials)
            except OneDriveOAuthInvalid as error:
                raise SourceRemoteUnauthorized() from error
            self.updated_encrypted_credentials = self.cipher.encrypt_credentials(credentials)
        return credentials

    def folders(self, *, encrypted_credentials: str | None) -> list[RemoteFolder]:
        return self.client.folders(credentials=self._credentials(encrypted_credentials))

    def discover(
        self, *, encrypted_credentials: str | None, selections: list[WorkspaceFolderSelection]
    ) -> DiscoveryResult:
        credentials = self._credentials(encrypted_credentials)
        changes: dict[str, dict[str, Any]] = {}
        removals: set[str] = set()
        links: dict[UUID, str | None] = {}
        root_id = (
            self.client.root_id(credentials=credentials)
            if any(item.kind == "root_files" for item in selections)
            else None
        )
        fresh_snapshot_required = False
        delta_expired = False
        if not fresh_snapshot_required:
            try:
                for selection in selections:
                    selection_changes: dict[str, dict[str, Any]] = {}
                    selection_removals: set[str] = set()
                    raw_cursor = (
                        self.cipher.decrypt_cursor(selection.encrypted_delta_link)
                        if selection.encrypted_delta_link
                        else None
                    )
                    page = self.client.delta(
                        credentials=credentials, selection=selection, cursor=raw_cursor
                    )
                    links[selection.id] = page.delta_link
                    for item in page.items:
                        if isinstance(item.get("folder"), dict):
                            # Delta may report a folder without its descendants. Re-enumerate
                            # non-root-file scopes so newly included subtrees are indexed and
                            # removed subtrees are reconciled.
                            if selection.kind != "root_files":
                                fresh_snapshot_required = True
                            continue
                        external_id = self.client._external_id(item, fallback_drive_id=None)
                        parent = item.get("parentReference") or {}
                        is_root_item = isinstance(parent, dict) and (
                            parent.get("id") == root_id
                            or self.client._parent_ids(parent) == ("root",)
                        )
                        if "deleted" in item or (
                            selection.kind == "root_files" and not is_root_item
                        ):
                            selection_changes.pop(external_id, None)
                            selection_removals.add(external_id)
                        else:
                            selection_removals.discard(external_id)
                            selection_changes[external_id] = item
                    changes.update(selection_changes)
                    removals.update(selection_removals)
                    # A file still present in any selected scope survives a
                    # tombstone from another overlapping selection.
                    removals.difference_update(changes)
            except (OneDriveCursorInvalid, OneDriveDeltaExpired):
                delta_expired = True
                fresh_snapshot_required = True
        if fresh_snapshot_required:
            files: dict[str, dict[str, Any]] = {}
            for selection in selections:
                for item in self.client.list_files(credentials=credentials, selection=selection):
                    files[self.client._external_id(item, fallback_drive_id=None)] = item
                links[selection.id] = None
            changes = files
            removals.clear()
            full_snapshot = True
        else:
            full_snapshot = all(not item.encrypted_delta_link for item in selections)
        if delta_expired:
            links = {selection.id: None for selection in selections}
        extracted = self._read_changed(credentials, changes)
        return DiscoveryResult(
            documents=extracted,
            removed_file_ids=tuple(sorted(removals)),
            delta_links=links,
            full_snapshot=full_snapshot,
        )

    def _read_changed(
        self, credentials: OneDriveCredentials, changes: dict[str, dict[str, Any]]
    ) -> list[DiscoveredDocument]:
        files = sorted(changes.values(), key=lambda item: str(item.get("id") or ""))
        if not files:
            return []
        with ThreadPoolExecutor(max_workers=min(MAX_PAGE_WORKERS, len(files))) as executor:
            return list(executor.map(lambda item: self._read_one(credentials, item), files))

    def _read_one(
        self, credentials: OneDriveCredentials, item: dict[str, Any]
    ) -> DiscoveredDocument:
        remote_id = self.client._external_id(item)
        mime_type = str((item.get("file") or {}).get("mimeType") or "application/octet-stream")
        parent_reference = item.get("parentReference") or {}
        modified = item.get("lastModifiedDateTime")
        base = {
            "external_file_id": remote_id,
            "name": str(item.get("name") or "Untitled"),
            "mime_type": mime_type,
            "source_url": str(item.get("webUrl") or ""),
            "modified_at": datetime.fromisoformat(str(modified)) if modified else None,
            "parent_ids": self.client._parent_ids(parent_reference),
        }
        if mime_type not in ELIGIBLE_MIME_TYPES:
            return DiscoveredDocument(**base)
        content = self.client.read_file(credentials=credentials, item_id=remote_id)
        from app.ingestion.google_drive import _extract_blocks

        try:
            blocks = _extract_blocks(mime_type, content)
        except (BadZipFile, PackageNotFoundError, PdfReadError, UnicodeDecodeError, ValueError):
            # Remote bodies and exception details are deliberately not persisted.
            return DiscoveredDocument(**base, error_code="text_extraction_failed")
        return DiscoveredDocument(
            **base, text="\n\n".join(block.text for block in blocks), blocks=tuple(blocks)
        )


class OneDriveConnectionService:
    def __init__(self, session: Session, cipher: OneDriveCipher, client: MicrosoftGraphClient):
        self.session, self.cipher, self.client = session, cipher, client

    def require_admin(self, *, scope: OrganizationScope, user_id: UUID) -> None:
        member = self.session.scalar(
            select(Membership).where(
                Membership.organization_id == scope.organization_id,
                Membership.user_id == user_id,
                Membership.is_active.is_(True),
                Membership.role.in_([MembershipRole.OWNER, MembershipRole.ADMIN]),
            )
        )
        if member is None:
            raise OneDriveAccessDenied("integration access denied")

    def begin(
        self,
        *,
        scope: OrganizationScope,
        user_id: UUID,
        session_secret: str,
        source_id: UUID | None = None,
    ) -> str:
        self.require_admin(scope=scope, user_id=user_id)
        self.cipher._fernet()
        if (
            source_id
            and self.session.scalar(
                select(DataSource).where(
                    DataSource.id == source_id,
                    DataSource.organization_id == scope.organization_id,
                    DataSource.provider == "onedrive",
                )
            )
            is None
        ):
            raise OneDriveAccessDenied("OneDrive source is invalid")
        raw_state = secrets.token_urlsafe(32)
        self.session.add(
            OAuthConnectionState(
                organization_id=scope.organization_id,
                user_id=user_id,
                source_id=source_id,
                session_hash=hash_secret(session_secret),
                state_hash=hash_secret(raw_state),
                expires_at=datetime.now(UTC) + timedelta(minutes=10),
            )
        )
        self.session.flush()
        return self.client.authorization_url(state=raw_state)

    def cancel(self, *, raw_state: str, session_secret: str) -> UUID:
        state = self.session.scalar(
            select(OAuthConnectionState)
            .where(
                OAuthConnectionState.state_hash == hash_secret(raw_state),
                OAuthConnectionState.consumed_at.is_(None),
                OAuthConnectionState.expires_at > datetime.now(UTC),
            )
            .with_for_update()
        )
        active_session = (
            self.session.scalar(
                select(UserSession).where(
                    UserSession.secret_hash == hash_secret(session_secret),
                    UserSession.user_id == state.user_id if state else False,
                    UserSession.revoked_at.is_(None),
                    UserSession.expires_at > datetime.now(UTC),
                )
            )
            if state
            else None
        )
        if (
            state is None
            or state.session_hash != hash_secret(session_secret)
            or active_session is None
        ):
            raise OneDriveOAuthInvalid("OAuth state is invalid")
        self.require_admin(scope=OrganizationScope(state.organization_id), user_id=state.user_id)
        state.consumed_at = datetime.now(UTC)
        self.session.flush()
        return state.organization_id

    def complete(self, *, raw_state: str, code: str, session_secret: str) -> DataSource:
        state = self.session.scalar(
            select(OAuthConnectionState)
            .where(
                OAuthConnectionState.state_hash == hash_secret(raw_state),
                OAuthConnectionState.consumed_at.is_(None),
                OAuthConnectionState.expires_at > datetime.now(UTC),
            )
            .with_for_update()
        )
        active_session = (
            self.session.scalar(
                select(UserSession).where(
                    UserSession.secret_hash == hash_secret(session_secret),
                    UserSession.user_id == state.user_id if state else False,
                    UserSession.revoked_at.is_(None),
                    UserSession.expires_at > datetime.now(UTC),
                )
            )
            if state
            else None
        )
        if (
            state is None
            or state.session_hash != hash_secret(session_secret)
            or active_session is None
        ):
            raise OneDriveOAuthInvalid("OAuth state is invalid")
        self.require_admin(scope=OrganizationScope(state.organization_id), user_id=state.user_id)
        credentials = self.client.exchange_code(code=code)
        account_email = self.client.account_email(credentials=credentials)
        drive_id = self.client.drive_id(credentials=credentials)
        source = (
            self.session.scalar(
                select(DataSource).where(
                    DataSource.id == state.source_id,
                    DataSource.organization_id == state.organization_id,
                    DataSource.provider == "onedrive",
                )
            )
            if state.source_id
            else self.session.scalar(
                select(DataSource)
                .where(
                    DataSource.organization_id == state.organization_id,
                    DataSource.provider == "onedrive",
                )
                .order_by(DataSource.created_at.desc(), DataSource.id.desc())
            )
        )
        if state.source_id is not None and source is None:
            raise OneDriveAccessDenied("OneDrive source is invalid")
        account_changed = source is not None and (
            (source.provider_account_id and source.provider_account_id != drive_id)
            or (
                not source.provider_account_id
                and source.account_email
                and source.account_email != account_email
            )
        )
        if account_changed:
            state.consumed_at = datetime.now(UTC)
            self.session.flush()
            raise OneDriveAccountMismatch(state.organization_id)
        if source is None:
            source = DataSource(
                organization_id=state.organization_id,
                provider="onedrive",
                encrypted_credentials=self.cipher.encrypt_credentials(credentials),
                status="connected",
                account_email=account_email,
                provider_account_id=drive_id,
                connected_by_user_id=state.user_id,
            )
            self.session.add(source)
        else:
            source.encrypted_credentials = self.cipher.encrypt_credentials(credentials)
            source.status = "connected"
            source.account_email = account_email
            source.provider_account_id = drive_id
            source.connected_by_user_id = state.user_id
            for selection in self.session.scalars(
                select(WorkspaceFolderSelection)
                .join(
                    WorkspaceFolder,
                    WorkspaceFolderSelection.workspace_folder_id == WorkspaceFolder.id,
                )
                .where(WorkspaceFolder.source_id == source.id)
            ):
                # Reestablish delta tokens after reauthorization; only this
                # source's bound drive identity is allowed to continue.
                selection.encrypted_delta_link = None
        state.consumed_at = datetime.now(UTC)
        self.session.flush()
        self.session.add(
            AuditLog(
                organization_id=state.organization_id,
                actor_user_id=state.user_id,
                action="data_source.connected",
                target_type="data_source",
                target_id=source.id,
            )
        )
        self.session.flush()
        return source

    def disconnect(self, *, scope: OrganizationScope, user_id: UUID, source_id: UUID) -> DataSource:
        self.require_admin(scope=scope, user_id=user_id)
        source = self.session.scalar(
            select(DataSource).where(
                DataSource.id == source_id,
                DataSource.organization_id == scope.organization_id,
                DataSource.provider == "onedrive",
            )
        )
        if source is None:
            raise OneDriveAccessDenied("OneDrive source is invalid")
        source.encrypted_credentials = None
        # Retain account identity so a later reconnect cannot silently bind
        # existing scopes and indexed documents to a different drive.
        source.status = "disconnected"
        self.session.execute(
            update(OAuthConnectionState)
            .where(
                OAuthConnectionState.source_id == source.id,
                OAuthConnectionState.consumed_at.is_(None),
            )
            .values(consumed_at=datetime.now(UTC))
        )
        for selection in self.session.scalars(
            select(WorkspaceFolderSelection)
            .join(
                WorkspaceFolder,
                WorkspaceFolderSelection.workspace_folder_id == WorkspaceFolder.id,
            )
            .where(WorkspaceFolder.source_id == source.id)
        ):
            selection.encrypted_delta_link = None
        self.session.add(
            AuditLog(
                organization_id=scope.organization_id,
                actor_user_id=user_id,
                action="data_source.disconnected",
                target_type="data_source",
                target_id=source.id,
            )
        )
        self.session.flush()
        return source

    def folders(
        self, *, scope: OrganizationScope, user_id: UUID, source_id: UUID
    ) -> list[RemoteFolder]:
        self.require_admin(scope=scope, user_id=user_id)
        source = self.session.scalar(
            select(DataSource).where(
                DataSource.id == source_id,
                DataSource.organization_id == scope.organization_id,
                DataSource.status == "connected",
                DataSource.provider == "onedrive",
            )
        )
        if source is None or source.encrypted_credentials is None:
            raise OneDriveOAuthInvalid("OneDrive source is invalid")
        provider = OneDriveDocumentProvider(self.client, self.cipher)
        try:
            result = provider.folders(encrypted_credentials=source.encrypted_credentials)
        except SourceRemoteUnauthorized as error:
            source.status = "reauth_required"
            self.session.flush()
            raise OneDriveReauthRequired("OneDrive reauthorization required") from error
        if provider.updated_encrypted_credentials:
            source.encrypted_credentials = provider.updated_encrypted_credentials
            self.session.flush()
        return result
